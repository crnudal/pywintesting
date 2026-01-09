import csv
import shutil
from pathlib import Path
from datetime import datetime

# =========================
# Inputs (edit these)
# =========================
ROOT_DIR = r"/path/to/folder_with_excels"
RECURSIVE = True

DRY_RUN = False              # True = no files saved, report still generated
MAKE_BACKUP = True           # True = copy original to *.bak_<timestamp>

REPORT_CSV = "format_string_replacements_report.csv"

FIND_TEXT = "yyyyMMdd-HH:mm:ss.SSSSSS"
REPLACE_TEXT = "yyyyMMdd-HH:mm:ss.SSSSSSSSS"

PROCESS_XLSX = True
PROCESS_XLS = True           # requires optional packages (see notes below)

# =========================
# XLSX handler (openpyxl)
# =========================
def process_xlsx(path: Path):
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=False)
    file_changed = False

    # Collect: (sheet_name, header_value, column_letter, num_occurrences_replaced)
    changes = []

    for ws in wb.worksheets:
        col_change_counts = {}  # col_idx -> replaced occurrences count (sum across cells)

        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v and FIND_TEXT in v:
                    # Replace substring even if there is text before/after
                    occurrences = v.count(FIND_TEXT)
                    new_v = v.replace(FIND_TEXT, REPLACE_TEXT)

                    if new_v != v:
                        cell.value = new_v
                        file_changed = True
                        col_change_counts[cell.column] = col_change_counts.get(cell.column, 0) + occurrences

        # Map changed columns to row-1 header value
        for col_idx, cnt in sorted(col_change_counts.items()):
            header_cell = ws.cell(row=1, column=col_idx).value
            header_value = str(header_cell) if header_cell is not None else ""
            col_letter = ws.cell(row=1, column=col_idx).column_letter
            changes.append((ws.title, header_value, col_letter, cnt))

    if file_changed and not DRY_RUN:
        if MAKE_BACKUP:
            ts = datetime.now().strftime("%Y%m%d_%H%M%S")
            backup_path = path.with_suffix(path.suffix + f".bak_{ts}")
            shutil.copy2(path, backup_path)
        wb.save(path)

    return file_changed, changes


# =========================
# XLS handler (xlrd + xlutils)
# =========================
def process_xls(path: Path):
    """
    Edits .xls in-place while trying to preserve formatting:
    - reads with xlrd
    - writes with xlutils.copy to an xlwt workbook

    Requires:
      pip install xlrd==1.2.0 xlwt xlutils
    """
    import xlrd
    from xlutils.copy import copy as xl_copy

    book = xlrd.open_workbook(path, formatting_info=True)
    wbook = xl_copy(book)

    file_changed = False
    changes = []

    for si in range(book.nsheets):
        sheet = book.sheet_by_index(si)
        wsheet = wbook.get_sheet(si)
        sheet_name = sheet.name

        col_change_counts = {}  # col_idx (0-based) -> replaced occurrences count

        for r in range(sheet.nrows):
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_TEXT:
                    v = cell.value
                    if v and FIND_TEXT in v:
                        occurrences = v.count(FIND_TEXT)
                        new_v = v.replace(FIND_TEXT, REPLACE_TEXT)

                        if new_v != v:
                            wsheet.write(r, c, new_v)
                            file_changed = True
                            col_change_counts[c] = col_change_counts.get(c, 0) + occurrences

        # header in first row is row 0
        for c, cnt in sorted(col_change_counts.items()):
            header_cell = sheet.cell(0, c)
            header_value = str(header_cell.value) if header_cell.value is not None else ""
            col_letter = _col_letter(c + 1)
            changes.append((sheet_name, header_value, col_letter, cnt))

    if file_changed and not DRY_RUN:
        if MAKE_BACKUP:
            ts = datetime.now().strftime("%Y%m%d_%H%M%S")
            backup_path = path.with_suffix(path.suffix + f".bak_{ts}")
            shutil.copy2(path, backup_path)
        wbook.save(str(path))

    return file_changed, changes


def _col_letter(n: int) -> str:
    """1-based to Excel column letters."""
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


# =========================
# Main
# =========================
def main():
    root = Path(ROOT_DIR)
    if not root.exists():
        raise SystemExit(f"ROOT_DIR does not exist: {root}")

    patterns = []
    if PROCESS_XLSX:
        patterns += ["*.xlsx", "*.xlsm"]
    if PROCESS_XLS:
        patterns += ["*.xls"]

    files = []
    if RECURSIVE:
        for pat in patterns:
            files.extend(root.rglob(pat))
    else:
        for pat in patterns:
            files.extend(root.glob(pat))

    files = sorted(set(files))

    report_rows = []
    total_files_changed = 0

    # Try-import for xls deps only if needed
    xls_available = True
    if PROCESS_XLS:
        try:
            import xlrd  # noqa: F401
            import xlutils  # noqa: F401
            import xlwt  # noqa: F401
        except Exception:
            xls_available = False

    for f in files:
        suffix = f.suffix.lower()
        try:
            if suffix in [".xlsx", ".xlsm"] and PROCESS_XLSX:
                changed, changes = process_xlsx(f)
            elif suffix == ".xls" and PROCESS_XLS and xls_available:
                changed, changes = process_xls(f)
            elif suffix == ".xls" and PROCESS_XLS and not xls_available:
                changed, changes = False, []
                print(f"[SKIP] {f} (missing xls deps: xlrd==1.2.0, xlwt, xlutils)")
            else:
                continue
        except Exception as e:
            print(f"[ERROR] {f}: {e}")
            continue

        if changed:
            total_files_changed += 1
            for sheet_name, header_value, col_letter, cnt in changes:
                report_rows.append({
                    "file": str(f),
                    "sheet": sheet_name,
                    "column_letter": col_letter,
                    "header_row1_value": header_value,
                    "occurrences_replaced": cnt,
                })

    # Write report
    if report_rows:
        with open(REPORT_CSV, "w", newline="", encoding="utf-8") as fp:
            writer = csv.DictWriter(fp, fieldnames=[
                "file", "sheet", "column_letter", "header_row1_value", "occurrences_replaced"
            ])
            writer.writeheader()
            writer.writerows(report_rows)

    print(f"Scanned files: {len(files)}")
    print(f"Files changed: {total_files_changed}")
    print(f"Report: {REPORT_CSV}")
    if DRY_RUN:
        print("DRY_RUN=True, no files were saved.")

if __name__ == "__main__":
    main()
